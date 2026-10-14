#!/usr/bin/env python3
"""Read ONE sheet of a manual QA checklist into JSON for conversion to the common template.

Source layout (e.g. "OEM B_L List_Checklist.xlsx"):
  No. | Testcase ID | Test Case Desc. | Precondition | Steps / Check point | Expected Result |
  Status | Priority | Remark
with section-header rows between cases (only the ID column filled; column D may hold a
precondition for the whole section). Section depth comes from the header row's fill colour
(darker = higher level); without fills, bold = level 1 and plain = level 2.

Emits:
  { "file", "sheet", "module", "hidden_rows", "group_level",
    "suggest": {"prefix", "repo_path", "screen"},
    "groups": [{"name", "abbr", "count"}],
    "cases": [{"key", "row", "hidden", "path", "group", "section_pre", "src_id", "desc",
               "pre", "steps", "expected", "priority", "status", "remark"}] }

`key` = "R<row>" (stable source-row key). `group` = the case's section at `group_level` (the
shallowest level with 2+ distinct names); TC # numbering in write_template.py restarts per group.

Usage:
  python extract_manual.py <file.xlsx> --list
  python extract_manual.py <file.xlsx> --sheet NAME [--chunks-dir DIR --chunk-size 40]
--chunks-dir also writes chunk_<k>.json (cases + sheet context) for parallel conversion.
"""
import argparse
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "testforge-gen" / "scripts"))
from read_cases import canon  # noqa: E402  (shared header normalizer)

ALIASES = {
    "no": ["no", "num", "number", "stt"],
    "id": ["testcaseid", "id", "tcid", "caseid", "testcase"],
    "desc": ["testcasedesc", "testcasedescription", "description", "title", "summary",
             "testscenario", "testcasename"],
    "pre": ["precondition", "preconditions", "pre"],
    "steps": ["stepscheckpoint", "steps", "step", "teststeps", "action", "actions",
              "stepcheckpoint"],
    "expected": ["expectedresult", "expectedresults", "expected", "expectedoutput"],
    "status": ["status"],
    "priority": ["priority"],
    "remark": ["remark", "remarks", "note", "notes"],
}
HEADER_SCAN = 10


def map_columns(headers):
    ch = [canon(h) for h in headers]
    col = {}
    for field, names in ALIASES.items():
        for name in names:  # alias order = preference
            if name in ch:
                col[field] = ch.index(name)
                break
    return col


def text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def load(path):
    try:
        import openpyxl
    except ImportError:
        sys.exit("openpyxl not installed. Run: python -m pip install openpyxl")
    # styles needed for section levels, so not read_only; data_only for formula results
    return openpyxl.load_workbook(path, data_only=True)


def find_header(ws):
    for r in range(1, min(ws.max_row, HEADER_SCAN) + 1):
        col = map_columns([text(c.value) for c in ws[r]])
        if "steps" in col and "desc" in col:
            return r, col
    return None, None


def style_key(cell):
    f = cell.fill
    if f is not None and f.fill_type == "solid":
        rgb = f.fgColor.rgb if isinstance(f.fgColor.rgb, str) else None
        if rgb and rgb[-6:].upper() not in ("FFFFFF", "000000"):
            return ("fill", rgb[-6:].upper())
    return ("bold", "") if cell.font and cell.font.b else ("plain", "")


def luminance(hex6):
    r, g, b = (int(hex6[i:i + 2], 16) for i in (0, 2, 4))
    return 0.299 * r + 0.587 * g + 0.114 * b


def rank_levels(keys):
    """style keys -> level (0 = top). Fills by darkness, then bold, then plain."""
    fills = sorted({k[1] for k in keys if k[0] == "fill"}, key=luminance)
    order = [("fill", h) for h in fills] + [("bold", ""), ("plain", "")]
    used = [k for k in order if k in keys]
    return {k: i for i, k in enumerate(used)}


WORD = re.compile(r"[A-Za-z0-9]+")
STOP = {"the", "a", "an", "as", "and", "or", "of", "in", "on", "to", "by", "for", "with"}


def abbreviate(name):
    """'UI/UX (Layout, ...)' -> UIUX, 'Permission' -> PERM, 'MB/L List grid' -> MBLLG."""
    base = re.sub(r"\(.*?\)|:.*$", " ", name)
    base = re.sub(r"(?<=\w)/(?=\w)", "", base)  # MB/L -> MBL, UI/UX -> UIUX
    base = re.sub(r"(?i)pop-up", "popup", base)
    words = WORD.findall(base)
    if len(words) > 1:
        words = [w for w in words if w.lower() not in STOP] or words
    if not words:
        return "SEC"
    if len(words) == 1:
        return words[0][:4].upper()
    out = "".join(w if (w.isupper() and len(w) > 1) else w[0] for w in words).upper()
    return out[:6]


def unique_abbrs(names):
    out, seen = {}, set()
    for n in names:
        a = abbreviate(n)
        cand, k = a, 2
        while cand in seen:
            cand, k = f"{a}{k}", k + 1
        seen.add(cand)
        out[n] = cand
    return out


def read_module(wb):
    """Last non-empty 'Module(s):' value of an Information sheet (tester section wins)."""
    for ws in wb.worksheets:
        if canon(ws.title) != "information":
            continue
        found = ""
        for row in ws.iter_rows():
            vals = [text(c.value) for c in row]
            for i, v in enumerate(vals):
                if canon(v) == "modules":
                    rest = [x for x in vals[i + 1:] if x]
                    if rest:
                        found = rest[0]
        return found
    return ""


def list_sheets(wb):
    out = []
    for ws in wb.worksheets:
        h, col = find_header(ws)
        if not h:
            out.append({"sheet": ws.title, "cases": 0, "note": "no Desc + Steps header"})
            continue
        n = sum(1 for r in range(h + 1, ws.max_row + 1)
                if text(ws.cell(r, col["desc"] + 1).value)
                or text(ws.cell(r, col["steps"] + 1).value))
        out.append({"sheet": ws.title, "cases": n})
    return out


def extract(wb, sheet):
    if sheet not in wb.sheetnames:
        sys.exit(f"Sheet not found: {sheet!r}. Sheets: {wb.sheetnames}")
    ws = wb[sheet]
    h, col = find_header(ws)
    if not h:
        sys.exit(f"Sheet {sheet!r} has no 'Test Case Desc.' + 'Steps / Check point' header "
                 f"in its first {HEADER_SCAN} rows.")
    id_col = col.get("id", col.get("no", 0))

    def cell(r, field):
        i = col.get(field)
        return text(ws.cell(r, i + 1).value) if i is not None else ""

    rows = []  # ("section", r, name, pre, style) | ("case", r)
    for r in range(h + 1, ws.max_row + 1):
        desc, steps = cell(r, "desc"), cell(r, "steps")
        if desc or steps:
            rows.append(("case", r))
            continue
        name = text(ws.cell(r, id_col + 1).value)
        if name:
            rows.append(("section", r, name, cell(r, "pre"), style_key(ws.cell(r, id_col + 1))))
    levels = rank_levels({x[4] for x in rows if x[0] == "section"})

    stack = []  # [(level, name, pre)]
    cases, hidden = [], 0
    for x in rows:
        if x[0] == "section":
            lvl = levels[x[4]]
            while stack and stack[-1][0] >= lvl:
                stack.pop()
            stack.append((lvl, x[2], x[3]))
            continue
        r = x[1]
        is_hidden = bool(ws.row_dimensions[r].hidden)
        hidden += is_hidden
        pres = [p for _, _, p in stack if p and not p.lower().startswith("note")]
        cases.append({
            "key": f"R{r}", "row": r, "hidden": is_hidden,
            "path": [(lvl, n) for lvl, n, _ in stack],
            "section_pre": "\n".join(pres),
            "src_id": cell(r, "id"), "desc": cell(r, "desc"), "pre": cell(r, "pre"),
            "steps": cell(r, "steps"), "expected": cell(r, "expected"),
            "priority": cell(r, "priority"), "status": cell(r, "status"),
            "remark": cell(r, "remark"),
        })

    # group level = shallowest level with 2+ distinct section names
    by_level = {}
    for c in cases:
        for lvl, n in c["path"]:
            by_level.setdefault(lvl, set()).add(n)
    group_level = next((l for l in sorted(by_level) if len(by_level[l]) > 1),
                       min(by_level) if by_level else 0)
    names = []
    for c in cases:
        at = [n for lvl, n in c["path"] if lvl <= group_level]
        c["group"] = at[-1] if at else sheet
        c["path"] = [n for _, n in c["path"]]
        if c["group"] not in names:
            names.append(c["group"])
    abbr = unique_abbrs(names)

    module = read_module(wb)
    screen = module.split(">")[-1].strip() if module else sheet
    top = [n for lvl, n in sorted(((l, n) for l, s in by_level.items() for n in s))
           if lvl == 0]
    if len(top) == 1 and canon(top[0]) not in ("mainui", canon(screen)):
        screen = top[0]  # a pop-up sheet: its own top section names the screen
    module_path = "/".join(p.strip() for p in module.split(">")) if module else ""
    if module_path and screen not in module_path:
        module_path = f"{module_path}/{screen}"
    return {
        "sheet": sheet, "module": module, "hidden_rows": hidden, "group_level": group_level,
        "suggest": {"prefix": abbreviate(screen), "screen": screen,
                    "repo_path": f"CARIS Upgrade/{module_path or screen}"},
        "groups": [{"name": n, "abbr": abbr[n],
                    "count": sum(1 for c in cases if c["group"] == n)} for n in names],
        "cases": cases,
    }


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("file")
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--sheet")
    g.add_argument("--list", action="store_true")
    ap.add_argument("--chunks-dir")
    ap.add_argument("--chunk-size", type=int, default=40)
    args = ap.parse_args()
    for stream in (sys.stdout, sys.stderr):  # Windows pipes default to cp1252
        stream.reconfigure(encoding="utf-8")

    path = Path(args.file)
    if not path.exists():
        sys.exit(f"File not found: {path}")
    wb = load(path)
    if args.list:
        print(json.dumps(list_sheets(wb), indent=2, ensure_ascii=False))
        return 0

    out = {"file": str(path), **extract(wb, args.sheet)}
    if args.chunks_dir:
        d = Path(args.chunks_dir)
        d.mkdir(parents=True, exist_ok=True)
        ctx = {k: out[k] for k in ("sheet", "module")} | {"screen": out["suggest"]["screen"]}
        size = max(1, args.chunk_size)
        for k in range(0, len(out["cases"]), size):
            chunk = {**ctx, "cases": out["cases"][k:k + size]}
            (d / f"chunk_{k // size + 1}.json").write_text(
                json.dumps(chunk, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(out, indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
