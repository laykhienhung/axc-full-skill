#!/usr/bin/env python3
"""Self-test for testforge-convert: extract_manual.py, write_template.py, and the round trip
through testforge-gen's read_cases.py.

Usage: python selftest_convert.py [extract|writer|roundtrip|all]   (default: all)
Builds throwaway workbooks in a temp dir; never touches the repo.
"""
import json
import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
EXTRACT = HERE / "extract_manual.py"
WRITER = HERE / "write_template.py"
READER = HERE.parent.parent / "testforge-gen" / "scripts" / "read_cases.py"
HEAD = ["No.", "Testcase ID", "Test Case Desc.", "Precondition", "Steps / Check point",
        "Expected Result", "Status", "Priority", "Remark"]
FILLS = {0: "3C78D8", 1: "6D9EEB", 2: "A4C2F4"}  # darker = higher level


def run(script, *args):
    p = subprocess.run([sys.executable, str(script), *map(str, args)],
                       capture_output=True, text=True, encoding="utf-8")
    return p.returncode, p.stdout, p.stderr


def check(cond, msg):
    if not cond:
        raise AssertionError(msg)
    print(f"  ok  {msg}")


def checklist(path):
    import openpyxl
    from openpyxl.styles import PatternFill
    wb = openpyxl.Workbook()
    info = wb.active
    info.title = "Information"
    info.append(["", "TEST REPORT BY DEV"])
    info.append([1, "Module(s):", None])
    info.append(["", "TEST REPORT BY TESTER"])
    info.append([4, "Module(s):", "Ocean Export > Master B/L > OEM B/L List"])
    ws = wb.create_sheet("Main UI")
    ws.append(HEAD)

    def section(name, level, pre=None):
        ws.append([None, name, None, pre])
        ws.cell(ws.max_row, 2).fill = PatternFill("solid", fgColor="FF" + FILLS[level])

    section("Main UI", 0)
    section("Permission", 1)
    ws.append([1, "_001", "Verify no access", "No permission", "1. Open screen", "Denied",
               "Passed", "High", ""])
    ws.append([2, "_002", "Verify access", "Has permission", "1. Open screen", "Shown",
               "Pending", "High", "not ready"])
    section("UI/UX (Layout, Label)", 1, "The user has permission")
    section("Menu Tree", 2)
    ws.append([3, "_001", "Menu name", "", "1. Open screen\n2. Look at menu", "- Spelling",
               "Failed", "Medium", "BUG-1"])
    ws.row_dimensions[ws.max_row].hidden = True
    section("Commodity Search pop-up", 2, "Note: http://example")
    ws.append([4, "_002", "Popup name", "", "1. Open popup", "- CSS", "", "", ""])
    wb.create_sheet("FORMULA").append(["MBL HAS", "x"])
    wb.save(path)


def extract(tmp):
    book = tmp / "checklist.xlsx"
    checklist(book)
    rc, out, err = run(EXTRACT, book, "--sheet", "Main UI", "--chunks-dir", tmp / "chunks",
                       "--chunk-size", 3)
    assert rc == 0, err
    (tmp / "extract.json").write_text(out, encoding="utf-8")
    return json.loads(out)


def converted(cases, bad=None):
    out = []
    for c in cases:
        out.append({"key": c["key"], "type": "Manual" if c["key"] == "R8" else "Auto",
                    "type_reason": "", "summary": f"Verify {c['desc']}",
                    "priority": c["priority"] or "Medium", "pre": "1. Logged in",
                    "action": "1. Open the screen\n2. Click [Search]",
                    "input": "1. Name = A×101 [suggested]\n2. Filing No. = TBD",
                    "checkpoint": "1. Screen opens\n2. Search result",
                    "expected": f"1. Screen is displayed\n2. {c['expected'] or 'Result shown'}"})
    if bad:
        bad(out)
    return out


def test_extract(tmp):
    print("extract")
    d = extract(tmp)
    rc, out, _ = run(EXTRACT, tmp / "checklist.xlsx", "--list")
    lst = {x["sheet"]: x["cases"] for x in json.loads(out)}
    check(lst == {"Information": 0, "Main UI": 4, "FORMULA": 0}, "--list counts cases per sheet")
    check(len(d["cases"]) == 4, "4 cases, section rows excluded")
    check(d["group_level"] == 1, "group level = first level with 2+ names")
    check([(g["name"], g["abbr"], g["count"]) for g in d["groups"]]
          == [("Permission", "PERM", 2), ("UI/UX (Layout, Label)", "UIUX", 2)],
          "groups + abbreviations in source order")
    c3 = d["cases"][2]
    check(c3["path"] == ["Main UI", "UI/UX (Layout, Label)", "Menu Tree"], "section path by fill")
    check(c3["section_pre"] == "The user has permission", "section precondition inherited")
    check(c3["hidden"] and d["hidden_rows"] == 1, "hidden rows flagged")
    check(d["cases"][3]["section_pre"] == "The user has permission",
          "'Note:' section text is not a precondition")
    check(d["cases"][0]["section_pre"] == "", "sibling section does not leak precondition")
    check(d["suggest"]["prefix"] == "OEMBLL" and d["suggest"]["screen"] == "OEM B/L List",
          "prefix/screen from Information Module(s)")
    check(d["suggest"]["repo_path"] == "CARIS Upgrade/Ocean Export/Master B/L/OEM B/L List",
          "repo path suggestion")
    chunks = sorted((tmp / "chunks").glob("chunk_*.json"))
    check([len(json.loads(p.read_text(encoding="utf-8"))["cases"]) for p in chunks] == [3, 1],
          "chunk files split by --chunk-size")
    rc, _, err = run(EXTRACT, tmp / "checklist.xlsx", "--sheet", "FORMULA")
    check(rc != 0 and "header" in err, "sheet without Desc+Steps header is rejected")


def write(tmp, conv, name, *extra):
    cj = tmp / f"{name}.json"
    cj.write_text(json.dumps(conv), encoding="utf-8")
    out = tmp / f"{name}.xlsx"
    return run(WRITER, tmp / "extract.json", cj, "--out", out, "--prefix", "oem-bll",
               "--repo-path", "CARIS Upgrade/X", *extra), out


def test_writer(tmp):
    print("writer")
    d = extract(tmp)
    cases = d["cases"]

    def misalign(o):
        o[1]["input"] = "1. -"

    def badtype(o):
        o[0]["type"] = "Semi"

    def drop(o):
        o.pop()

    def dup(o):
        o.append(dict(o[0]))

    def suggested_outside_input(o):
        o[0]["expected"] = "1. Screen is displayed\n2. No Data Found [suggested]"

    for label, fn, key in [("misaligned step counts", misalign, "R5"),
                           ("bad type", badtype, "R4"),
                           ("[suggested] outside Input Data", suggested_outside_input, "R4"),
                           ("missing key", drop, "R10"),
                           ("duplicate key", dup, "R4")]:
        (rc, _, err), out = write(tmp, converted(cases, fn), f"bad_{label.replace(' ', '_')}")
        check(rc == 2 and f'"{key}"' in err and not out.exists(), f"rejects {label}")

    def renumber(o):
        o[0]["action"] = "1. A\n3. B"
        o[0]["input"] = o[0]["checkpoint"] = o[0]["expected"] = "1. -\n3. -"

    (rc, _, err), _ = write(tmp, converted(cases, renumber), "bad_numbering")
    check(rc == 2 and "1..n" in err, "rejects gaps in numbering")

    (rc, out, err), xlsx = write(tmp, converted(cases), "good")
    assert rc == 0, err
    s = json.loads(out)
    check(s["cases"] == 4 and s["by_type"] == {"Auto": 3, "Manual": 1}, "summary counts")
    check(s["tc_range"] == ["OEMBLL_PERM_001", "OEMBLL_UIUX_002"],
          "TC # = PREFIX_ABBR_NNN, restarting per group")
    check(s["input_values"] == {"suggested": 4, "tbd": 4}, "suggested / TBD input counts")
    (rc, _, err), _ = write(tmp, converted(cases), "good")
    check(rc != 0 and "--force" in err, "refuses to overwrite without --force")

    abbr = tmp / "abbr.json"
    abbr.write_text(json.dumps({"UI/UX (Layout, Label)": "UI"}), encoding="utf-8")
    (rc, out, _), _ = write(tmp, converted(cases), "with_abbr", "--abbr", abbr)
    check(json.loads(out)["tc_range"][-1] == "OEMBLL_UI_002", "--abbr overrides a group")

    import openpyxl
    ws = openpyxl.load_workbook(xlsx).worksheets[0]
    check(ws.title == "OEM BL List", "sheet named after screen, '/' stripped")
    check(ws.cell(1, 1).value == "TC #" and ws.max_column == 11, "template header kept (11 cols)")
    check(ws.cell(2, 6).value == "General" and ws.cell(2, 4).value == "CARIS Upgrade/X",
          "regression + repo path filled")
    check(ws.cell(2, 8).alignment.wrap_text and ws.cell(2, 1).font.name == "Arial",
          "data rows use the template's data style")


def test_roundtrip(tmp):
    print("roundtrip")
    d = extract(tmp)
    (rc, _, err), xlsx = write(tmp, converted(d["cases"]), "rt")
    assert rc == 0, err
    rc, out, err = run(READER, xlsx)
    assert rc == 0, err
    cases = json.loads(out)
    check([c["id"] for c in cases] == ["OEMBLL_PERM_001", "OEMBLL_PERM_002",
                                       "OEMBLL_UIUX_001", "OEMBLL_UIUX_002"],
          "read_cases maps 'TC #' to id")
    c = cases[0]
    check(c["title"] == "Verify Verify no access", "Summary -> title")
    check([s["action"] for s in c["steps"]] == ["Open the screen", "Click [Search]"],
          "Action split into steps")
    check([s["expected"] for s in c["steps"]] == ["Screen is displayed", "Denied"],
          "Expected aligned per step")
    check(c["preconditions"] == "1. Logged in" and "Filing No. = TBD" in c["data"],
          "Pre_Condition + Input Data carried")
    check("Name = A×101 [suggested]" in c["data"], "[suggested] tag survives into data")


TESTS = {"extract": test_extract, "writer": test_writer, "roundtrip": test_roundtrip}


def main() -> int:
    for stream in (sys.stdout, sys.stderr):
        stream.reconfigure(encoding="utf-8")
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        sys.exit("openpyxl not installed. Run: python -m pip install openpyxl")
    which = sys.argv[1] if len(sys.argv) > 1 else "all"
    names = list(TESTS) if which == "all" else [which]
    for n in names:
        with tempfile.TemporaryDirectory() as t:
            TESTS[n](Path(t))
    print("ALL PASSED")
    return 0


if __name__ == "__main__":
    sys.exit(main())
